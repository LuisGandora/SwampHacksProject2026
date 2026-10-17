"""Economic data for the sites' companies, all from EIA (public, no key):

  data/plant_utilities.csv     plant -> operating utility (EIA utility ID) and balancing authority (grid operator)
  data/utility_monthly_sales.csv   EIA-861M retail sales per utility per month: revenue, MWh, customers
  data/ba_daily_demand.csv     EIA-930 daily electricity demand per balancing authority (weekly totals come from this)

Utilities don't publish weekly sales; EIA-930 demand (MWh served) is the closest weekly measure of sales volume.
Run from the project root:  python scraper/economics.py
"""
import collections
import csv
import datetime as dt
from pathlib import Path

import openpyxl

from relate_weather import cached
from scrape_power_plants import DATA, RAW, download_eia860m

EIA861M = "https://www.eia.gov/electricity/data/eia861m"
EIA930 = "https://www.eia.gov/electricity/gridmonitor/knownissues/xls"
YEARS = range(2021, dt.date.today().year + 1)
MIN_SITES_PER_BA = 20  # only download grid-operator files (~75 MB each) that cover a meaningful number of sites


def plant_links(site_ids):
    rows = openpyxl.load_workbook(download_eia860m(), read_only=True)["Operating"].iter_rows(min_row=3, values_only=True)
    h = {k: i for i, k in enumerate(next(rows))}
    links = {}
    for r in rows:
        pid = str(r[h["Plant ID"]])
        if pid in site_ids and pid not in links:
            links[pid] = (r[h["Entity ID"]], r[h["Entity Name"]], r[h["Balancing Authority Code"]] or "")
    with open(DATA / "plant_utilities.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["plant_id", "utility_id", "utility_name", "balancing_authority"])
        w.writerows([pid, *v] for pid, v in sorted(links.items()))
    return links


def utility_sales(utility_ids):
    out = []
    for y in YEARS:
        name = f"sales_ult_cust_{y}.xlsx"
        url = f"{EIA861M}/xls/{name}" if y == dt.date.today().year else f"{EIA861M}/archive/xls/{name}"
        try:
            path = cached(name, url, headers={"User-Agent": "Mozilla/5.0"})
        except Exception as e:  # current-year file can lag; keep going with what exists
            print(f"  {name}: {e}")
            continue
        rows = openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True)
        group, _, header = next(rows), next(rows), next(rows)
        total = next(i for i, g in enumerate(group) if g and str(g).strip().upper() == "TOTAL")
        col = {k: i for i, k in enumerate(header)}
        for r in rows:
            if r[col["Utility Number"]] in utility_ids and isinstance(r[col["Year"]], int):
                out.append((r[col["Year"]], r[col["Month"]], r[col["Utility Number"]], r[col["Utility Name"]],
                            r[col["State"]], r[total], r[total + 1], r[total + 2]))
    with open(DATA / "utility_monthly_sales.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["year", "month", "utility_id", "utility_name", "state", "revenue_thousand_usd", "sales_mwh",
                    "customers"])
        w.writerows(out)
    return len(out)


def ba_demand(bas):
    out = []
    for ba in bas:
        path = cached(f"eia930_{ba}.xlsx", f"{EIA930}/{ba}.xlsx", headers={"User-Agent": "Mozilla/5.0"})
        rows = openpyxl.load_workbook(path, read_only=True)["Published Daily Data"].iter_rows(values_only=True)
        col = {k: i for i, k in enumerate(next(rows))}
        for r in rows:
            day, demand = r[col["Local date"]], r[col["Demand"]]
            if isinstance(day, dt.datetime) and isinstance(demand, (int, float)) and day.year >= min(YEARS):
                out.append((ba, day.date().isoformat(), demand))
        print(f"  {ba}: {sum(o[0] == ba for o in out)} days")
    with open(DATA / "ba_daily_demand.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["balancing_authority", "date", "demand_mwh"])
        w.writerows(out)
    return len(out)


def main():
    RAW.mkdir(parents=True, exist_ok=True)
    sites = {r["eia_plant_id"] for r in csv.DictReader(open(DATA / "power_plants.csv", encoding="utf-8"))}
    links = plant_links(sites)
    print(f"{len(links)} plants linked to utilities and grid operators")
    n = utility_sales({v[0] for v in links.values()})
    print(f"utility_monthly_sales: {n} utility-months")
    count = collections.Counter(v[2] for v in links.values() if v[2])
    bas = sorted(b for b, c in count.items() if c >= MIN_SITES_PER_BA)
    print(f"grid operators with >= {MIN_SITES_PER_BA} sites: {dict((b, count[b]) for b in bas)}")
    print(f"ba_daily_demand: {ba_demand(bas)} days")


if __name__ == "__main__":
    main()
